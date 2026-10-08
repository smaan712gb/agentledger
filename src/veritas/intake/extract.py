"""Turn any incoming file into parts with text (all open-source extractors).

PDF -> pypdf; scanned PDF/images -> Tesseract if installed, else the local vision
model; DOCX -> python-docx; XLSX -> openpyxl; CSV/TXT/JSON/XML/HTML/EDI -> decode;
EML -> body + attachments (recursively); ZIP -> members (recursively).
"""

from __future__ import annotations

import email
import email.policy
import io
import mimetypes
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from ..regwatch.documents import html_to_text

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp",
               ".tif": "image/tiff", ".tiff": "image/tiff", ".bmp": "image/bmp", ".heic": "image/heic"}
TEXT_EXT = {".txt", ".csv", ".tsv", ".json", ".xml", ".md", ".edi", ".835", ".x12", ".ofx", ".qfx", ".qif", ".iif", ".beancount"}
MAX_DEPTH = 4


@dataclass
class Part:
    name: str
    data: bytes
    media_type: str
    text: str = ""
    images: list[bytes] = field(default_factory=list)  # for vision when no text layer
    sender: str | None = None
    subject: str | None = None
    parent: str | None = None
    note: str = ""


def media_type(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    if data[:4] == b"%PDF":
        return "application/pdf"
    if data[:2] == b"PK" and ext not in (".docx", ".xlsx"):
        return "application/zip"
    if ext in IMAGE_TYPES:
        return IMAGE_TYPES[ext]
    if ext == ".eml" or (data[:200].lower().find(b"\nsubject:") >= 0 and data[:200].lower().find(b"from:") >= 0):
        return "message/rfc822"
    return mimetypes.guess_type(name)[0] or "application/octet-stream"


def explode(name: str, data: bytes, depth: int = 0, parent: str | None = None, sender: str | None = None,
            subject: str | None = None) -> list[Part]:
    """Split containers (zip, email) into leaf parts, extracting text from each."""
    mt = media_type(name, data)
    if depth > MAX_DEPTH:
        return [Part(name, data, mt, note="nesting too deep; not expanded", parent=parent)]
    if mt == "application/zip":
        out: list[Part] = []
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                if info.is_dir() or info.file_size > 50_000_000 or info.filename.startswith("__MACOSX"):
                    continue
                out += explode(Path(info.filename).name, z.read(info), depth + 1, name, sender, subject)
        return out
    if mt == "message/rfc822":
        msg = email.message_from_bytes(data, policy=email.policy.default)
        frm = email.utils.parseaddr(str(msg.get("from", "")))[1].lower() or sender
        subj = str(msg.get("subject", "")) or subject
        out = []
        body = msg.get_body(preferencelist=("plain", "html"))
        if body is not None:
            content = body.get_content()
            text = html_to_text(content) if body.get_content_type() == "text/html" else content
            if text.strip():
                out.append(Part(f"{Path(name).stem}-body.txt", text.encode(), "text/plain", text=text, sender=frm,
                                subject=subj, parent=name))
        for att in msg.iter_attachments():
            payload = att.get_payload(decode=True) or b""
            out += explode(att.get_filename() or "attachment.bin", payload, depth + 1, name, frm, subj)
        return out
    part = Part(name, data, mt, parent=parent, sender=sender, subject=subject)
    extract(part)
    return [part]


def extract(part: Part) -> None:
    ext = Path(part.name).suffix.lower()
    try:
        if part.media_type == "application/pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(part.data))
            part.text = "\n".join((p.extract_text() or "") for p in reader.pages[:50])
            if len(part.text.strip()) < 40:  # scanned: no text layer
                for page in reader.pages[:3]:
                    part.images += [img.data for img in page.images][:2]
                part.text = ocr_many(part.images) or ""
        elif part.media_type.startswith("image/"):
            part.images = [part.data]
            part.text = ocr_many(part.images) or ""
        elif ext == ".docx":
            import docx

            d = docx.Document(io.BytesIO(part.data))
            lines = [p.text for p in d.paragraphs]
            for t in d.tables:
                lines += [" | ".join(c.text for c in row.cells) for row in t.rows]
            part.text = "\n".join(lines)
        elif ext in (".xlsx", ".xlsm"):
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(part.data), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets[:10]:
                lines.append(f"# sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True, max_row=2000):
                    if any(v is not None for v in row):
                        lines.append(",".join("" if v is None else str(v) for v in row))
            part.text = "\n".join(lines)
        elif ext in (".html", ".htm"):
            part.text = html_to_text(part.data.decode("utf-8", "replace"))
        elif ext in TEXT_EXT or part.media_type.startswith("text/"):
            part.text = part.data.decode("utf-8", "replace")
        else:
            part.note = f"no extractor for {ext or part.media_type}"
    except Exception as e:
        part.note = f"extraction failed: {type(e).__name__}: {e}"


def ocr_many(images: list[bytes]) -> str | None:
    """Tesseract (open source) when installed; otherwise the caller falls back to the vision model."""
    exe = shutil.which("tesseract")
    if not exe or not images:
        return None
    texts = []
    with tempfile.TemporaryDirectory() as tmp:
        for i, img in enumerate(images):
            p = Path(tmp) / f"img{i}"
            p.write_bytes(img)
            r = subprocess.run([exe, str(p), "stdout"], capture_output=True, timeout=120)
            texts.append(r.stdout.decode("utf-8", "replace"))
    return "\n".join(texts)
